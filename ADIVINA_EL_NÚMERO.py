import lib_adivina
import openpyxl
import getpass
import tkinter as tk
from tkinter import filedialog
from openpyxl import load_workbook
import matplotlib.pyplot as plt

lib_adivina.welcome()
print('Hola! Bienvenido a Adivina el Numero. Antes de comenzar a jugar hagamos un pequeño set up.')
print('Elige dónde guardar y nombra el archivo: Estadisticas.xlsx, este archivo guardará los resultados del juego.\nSi el archivo ya existe, haz clic en él y selecciona si para reemplazar')

root = tk.Tk()
root.withdraw()
    
ubicacion_archivo = filedialog.asksaveasfilename(defaultextension=".xlsx", 
                                                     filetypes=[("Excel files", "*.xlsx")],
                                                     title="Elige dónde guardar Estadisticas.xlsx")
    
if not ubicacion_archivo:  
    print("No se seleccionó ningún archivo. Terminando el programa.")
    exit()

try:
    excelDocument = load_workbook(ubicacion_archivo)
    hojaSolitario = excelDocument["Solitario"]
    hoja2Jugadores = excelDocument["2Jugadores"]

except FileNotFoundError:
    excelDocument = openpyxl.Workbook()
    hojaSolitario = excelDocument.active
    hojaSolitario.title = "Solitario"
    hojaSolitario.append(["Nombre", "Nivel", "Ganadas"])

    hoja2Jugadores = excelDocument.create_sheet(title="2Jugadores")
    hoja2Jugadores.append(["Nombre", "Nivel", "Ganadas"])

excelDocument.save(ubicacion_archivo)
print("⭐️⭐️⭐️¡Archivo guardado exitosamente!⭐️⭐️⭐️")


def menu():
    lib_adivina.Menu()
    print("Selecciona una opcion: ")
    print("1. Partida modo solitario")
    print("2. Partida 2 Jugadores")
    print("3. Estadística")
    print("4. Salir")

def validaInt(mensaje=''):
    while True:
        try:
            opcion = int(input(mensaje))
            return opcion
        except ValueError:
            print("Error, ingresa un número válido.")

def valida(min, max):
    if min > max:
        print("Error, min no puede ser mas grande que max")
        return None

    while True:
        opcion = validaInt()
        if min <= opcion <= max:
            return opcion
        else:
            print(f'Error, escribe una opción entre {min} y {max}: ')

def menu2():
    lib_adivina.nivel()
    print("Selecciona un nivel: ")
    print("1. Fácil (20 intentos)")
    print("2. Medio (12 intentos)")
    print("3. Difícil (5 intentos)")
    print("4. Crea tu propio nivel personalizado")
    print("5. Regresar al menu principal")

def jugarSolitario (chances, min, max):
    if chances == 20 and min == 1 and max == 1000:
        nivel= "facil"
    elif chances == 12 and min == 1 and max == 1000:
        nivel= "medio"
    elif chances == 5 and min == 1 and max == 1000:
        nivel= "difícil"
    else:
        nivel = "personalizado"
    win=False
    numIntentado = []
    soliNumRandom=lib_adivina.giveMeNumber(min,max)
    print(f'Adivina un numero del {min} al {max}')
    while not win and chances>0:
        if len(numIntentado) == 0:
            print("Escribe tu primer numero")
        else:
            print("Estos son los numeros que has intrucido hasta el momento: " + str(numIntentado))
        adivinaSoliNum = validaInt()
        if adivinaSoliNum == soliNumRandom:
            win=True
        if adivinaSoliNum != soliNumRandom:
            chances = chances-1
            numIntentado.append(adivinaSoliNum)
            if adivinaSoliNum > soliNumRandom:
                print("El numero a adivinar es menor al intrucido")
            else:
                 print("El numero a adivinar es mayor al intrucido")
    if win:
        lib_adivina.win()
        name= input("Escribe tu nombre porfavor: ")
        hojaSolitario.append([name, nivel, 1]) 
    else:
        lib_adivina.lose()
        print("El numero era: " + str(soliNumRandom))
        name= input("Escribe tu nombre porfavor: ")
        hojaSolitario.append([name, nivel, 0])
    excelDocument.save(ubicacion_archivo)
      
def jugarSolitarioPersonalizadoNivel():
    chances= validaInt("Elije el numero de intentos: ")
    print("Ahora escoge el rango")
    min = validaInt("Elije el numero menor: ")
    max = validaInt("Elije el numero mayor: ")
    while min > max:
        print("Error el numero menor no puede ser mas grande que el mayor")
        min = validaInt("Elije el numero menor: ")
        max = validaInt("Elije el numero mayor: ")

    jugarSolitario (chances,min,max)

def jugar2Jugadores(chances, min, max, chancesPersonalizado=False):
    if chancesPersonalizado:
        chances = validaInt("Primero escribe el número de intentos para jugador 2: ")
        min = validaInt("Elije el número menor: ")
        max = validaInt("Elije el número mayor: ")
        while min > max:
            print("Error, el número menor no puede ser mayor al mayor")
            min = validaInt("Elije el número menor: ")
            max = validaInt("Elije el número mayor: ")

    if chances == 20 and min == 1 and max == 1000:
        nivel= "facil"
    elif chances == 12 and min == 1 and max == 1000:
        nivel= "medio"
    elif chances == 5 and min == 1 and max == 1000:
        nivel= "dificil"
    else:
        nivel = "personalizado"
    win = False
    numIntentado = []
    print(f"Bienvenidos a la partida para 2 jugadores.\nJugador número 1 escribirá un número entre {min} y {max}, jugador número 2 intentará adivinarlo.")
    NumJugador1 = int(getpass.getpass("Jugador número 1, introduce el número a adivinar (no se mostrará en pantalla): "))

    while NumJugador1 < min or NumJugador1 > max:
        NumJugador1 = int(getpass.getpass("Error, escribe un número dentro del rango seleccionado: "))

    while not win and chances>0:
        if len(numIntentado) == 0:
            print("Jugador 2 escribe tu primer numero: ")
        else:
            print("Estos son los numeros que has intrucido hasta el momento: " + str(numIntentado))
        NumJugador2 = validaInt()  
        if NumJugador2 == NumJugador1:
            win=True
        if NumJugador2 != NumJugador1:
            chances = chances-1
            numIntentado.append(NumJugador2)
            if NumJugador2 > NumJugador1:
                print("El numero a adivinar es menor al intrucido")
            else:
                 print("El numero a adivinar es mayor al intrucido")
    if win:
        lib_adivina.win()
        name= input("Jugador 2 Escribe tu nombre porfavor: ")
        hoja2Jugadores.append([name, nivel, 1])
    else:
        lib_adivina.lose()
        print("El numero era: " + str(NumJugador1))
        name= input("Jugador 2 Escribe tu nombre porfavor: ")
        hoja2Jugadores.append([name, nivel, 0])

    excelDocument.save(ubicacion_archivo)

def menuEstadistica():
    lib_adivina.estadistica()
    print("Selecciona una opcion: ")
    print("1. Estadísticas por usuario")
    print("2. Estadísticas por modo de juego")
    print("3. Estadísticas por dificultad")
    print("4. Hoja the Excel")
    print("5. Salir")

def nombreEstandar(nombre):
    return ' '.join(nombre.split()).lower() if isinstance(nombre, str) else nombre

def ganadasPorUsuario(hoja, usuario):
    total_ganadas = 0  
    usuario = nombreEstandar(usuario)
    for row in hoja.iter_rows(min_row=2, values_only=True):
        nombre, nivel, ganadas = row
        if nombreEstandar(nombre) == usuario:
            total_ganadas += ganadas  
    return total_ganadas

def perdidasPorUsuario(hoja, usuario):
    total_perdidas = 0  
    usuario = nombreEstandar(usuario)
    for row in hoja.iter_rows(min_row=2, values_only=True):
        nombre, nivel, ganadas = row
        if nombreEstandar(nombre) == usuario and ganadas == 0:
            total_perdidas += 1
    return total_perdidas

while True:
    opcionMenu1 = 0
    menu()
    opcionMenu1 = valida(1, 4)

    if opcionMenu1 == 4:
        break

    if opcionMenu1 == 1:
        menu2()
        opcionMenu2 = valida(1, 5)
        if opcionMenu2 == 1:
            jugarSolitario (20,1,1000)
        if opcionMenu2 == 2:
            jugarSolitario (12,1,1000)
        if opcionMenu2 == 3:
            jugarSolitario (5,1,1000)
        if opcionMenu2 == 4:
            jugarSolitarioPersonalizadoNivel()
            
    if opcionMenu1 == 2:
        menu2()
        opcionMenu2 = valida(1,5)
        if opcionMenu2 == 1:
            jugar2Jugadores (20,1,1000)
        if opcionMenu2 == 2:
            jugar2Jugadores (12,1,1000)
        if opcionMenu2 == 3:
            jugar2Jugadores (5,1,1000)
        if opcionMenu2 == 4:
            jugar2Jugadores(0, min, max, chancesPersonalizado=True)

    if opcionMenu1 == 3:
        menuEstadistica()
        opcionEstadistica = valida(1,6)

        if opcionEstadistica == 1:
                    usuario = input('Escribe el nombre del usuario que deseas buscar: ')
                    usuario= nombreEstandar(usuario)

                    for row in hojaSolitario.iter_rows(min_row=2, min_col=1, max_col=1):
                        cell = row[0]
                        cell.value = nombreEstandar(cell.value)
                        excelDocument.save(ubicacion_archivo)
                        
                    for row in hoja2Jugadores.iter_rows(min_row=2, min_col=1, max_col=1): 
                        cell = row[0]
                        cell.value = nombreEstandar(cell.value)
                        excelDocument.save(ubicacion_archivo)

                    usuarioGanadasSolitario = 0
                    usuarioGanadas2Jugadores = 0
                    usuarioPerdidasSolitario = 0
                    usuarioPerdidas2Jugadores = 0
                    usuarioGanadasSolitario = ganadasPorUsuario(hojaSolitario, usuario)
                    usuarioGanadas2Jugadores = ganadasPorUsuario(hoja2Jugadores, usuario)
                    usuarioPerdidasSolitario = perdidasPorUsuario(hojaSolitario, usuario)
                    usuarioPerdidas2Jugadores = perdidasPorUsuario(hoja2Jugadores, usuario)
                
                    labels = ['Ganadas', 'Perdidas']
                    usuarioSolitarioData = [usuarioGanadasSolitario, usuarioPerdidasSolitario]
                    usuario2JugadoresData = [usuarioGanadas2Jugadores, usuarioPerdidas2Jugadores]
                    plt.figure(figsize=(15, 5))
                    plt.subplot(1, 2, 1)
                    plt.bar(labels, usuarioSolitarioData, color=['green', 'red'])
                    plt.title(' Partidas jugadas en Solitario')
                    plt.ylabel(f'Ganadas y perdidas de {usuario}')

                    plt.subplot(1, 2, 2)
                    plt.bar(labels, usuario2JugadoresData, color=['green', 'red'])
                    plt.title('Partidas jugadas con 2 Jugadores')

                    plt.suptitle(f'Cantidad de Victorias Y Perdidas de {usuario} en Cada Modo de Juego')
                    plt.show()
        
        if opcionEstadistica == 2:
            solitarioGanadas = 0
            solitarioPerdidas = 0
            jugadores2Ganadas= 0
            jugadores2Perdidas= 0
            
            for row in hojaSolitario.iter_rows(min_row=2, values_only=True): 
                nombre, nivel, ganadas = row
                if ganadas == 1:
                    solitarioGanadas += 1
                elif ganadas == 0:
                    solitarioPerdidas += 1

            for row in hoja2Jugadores.iter_rows(min_row=2, values_only=True): 
                nombre, nivel, ganadas = row
                if ganadas == 1:
                    jugadores2Ganadas += 1
                elif ganadas == 0:
                    jugadores2Perdidas += 1
                    
            labels = ['Ganadas', 'Perdidas']
            solitario_data = [solitarioGanadas, solitarioPerdidas]
            jugadores_data = [jugadores2Ganadas, jugadores2Perdidas]
            plt.figure(figsize=(10, 5))
            plt.subplot(1, 2, 1)
            plt.bar(labels, solitario_data, color=['green', 'red'])
            plt.title('Partidas Solitario')
            plt.ylabel('Número de Personas')

            plt.subplot(1, 2, 2)
            plt.bar(labels, jugadores_data, color=['green', 'red'])
            plt.title('Partidas 2 Jugadores')

            plt.suptitle('Resultados de Partidas: Ganadas y Perdidas')
            plt.show()
        
        if opcionEstadistica == 3:
            facilSolitario = 0
            medioSolitario= 0
            dificillSolitario= 0
            personalizadoSolitario= 0
            facil2Jugadores = 0
            medio2Jugadores= 0
            dificil2Jugadores= 0
            personalizado2Jugadores= 0

            for row in hojaSolitario.iter_rows(min_row=2, values_only=True): 
                nombre, nivel, ganadas = row
                if nivel == "facil" and ganadas == 1:
                    facilSolitario += 1
                elif nivel == "medio" and ganadas == 1:
                    medioSolitario += 1
                elif nivel == "dificil" and ganadas == 1:
                     dificillSolitario+= 1
                elif nivel == "personalizado" and ganadas == 1:
                    personalizadoSolitario += 1

            for row in hoja2Jugadores.iter_rows(min_row=2, values_only=True): 
                nombre, nivel, ganadas = row
                if nivel == "facil" and ganadas == 1:
                    facil2Jugadores += 1
                elif nivel == "medio" and ganadas == 1:
                    medio2Jugadores += 1
                elif nivel == "dificil" and ganadas == 1:
                     dificil2Jugadores += 1
                elif nivel == "personalizado" and ganadas == 1:
                    personalizado2Jugadores += 1
            
            labels = ['Fácil', 'Medio', 'Difícil', 'Personalizado']
            nivelesSolitarioData = [facilSolitario, medioSolitario, dificillSolitario, personalizadoSolitario]
            nivel2jugadoresData = [facil2Jugadores, medio2Jugadores,dificil2Jugadores,personalizado2Jugadores]
            plt.figure(figsize=(20, 5))
            plt.subplot(1, 4, 1)
            plt.bar(labels, nivelesSolitarioData, color=['green', 'blue', 'red', 'orange'])
            plt.title(' Nivel de dificultad en Solitario')
            plt.ylabel('Número de Personas')

            plt.subplot(1, 4, 2)
            plt.bar(labels, nivel2jugadoresData, color=['green', 'blue', 'red', 'orange'])
            plt.title('Nivel de dificultad con 2 Jugadores')

            plt.suptitle('Partidas ganadas segun su nivel de dificultad')
            plt.show()

        if opcionEstadistica == 4:
            print("Cual hoja te gustaria ver? \n 1. Datos juego solitario  \n 2. Datos partida 2 jugadores")
            fichero= validaInt()
            while fichero < 1 or fichero > 2:
                fichero = validaInt("Error, escribe una opción entre 1 y 2: ")
            if fichero == 1:
                for row in hojaSolitario:
                    for cell in row:
                        print(str(cell.value) + " ", end="")  
                    print()
            elif fichero == 2:
                for row in hoja2Jugadores:
                    for cell in row:
                        print(str(cell.value) + " ", end="")  
                    print()
